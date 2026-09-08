# -*- coding: utf-8 -*-
"""Prueba de punta a punta del agente de productos.

Levanta la app real, la maneja con un navegador real y le pone detrás un Odoo
en memoria. Comprueba lo que los tests unitarios no pueden: que la pantalla,
la API y el motor se entiendan entre sí.

No la corre pytest (no empieza con test_): necesita Playwright y Chromium.

    pip install playwright && playwright install chromium
    python tests/e2e_productos.py

Si Chromium ya está en el sistema, indicar la ruta con CHROMIUM=/ruta/al/chrome.
"""
import os, sys, tempfile, threading, time, socket
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
os.chdir(RAIZ)

from unittest.mock import patch
import uvicorn
from playwright.sync_api import sync_playwright
from tests.test_productos import base
from productos import api as productos_api
import main

o, d = base()
IDENTIDAD = {"email": "lucas@kaironsrl.com.ar", "sub": "e2e"}

puerto = socket.socket(); puerto.bind(("127.0.0.1", 0)); PORT = puerto.getsockname()[1]; puerto.close()

parches = [patch.object(main, "verificar_token_access", return_value=IDENTIDAD),
           patch.object(productos_api, "_odoo", return_value=o),
           # /status informa disponibilidad mirando las variables de entorno,
           # no la conexion: son las que tendria Railway.
           patch.dict(os.environ, {"PRODUCTOS_ALLOWED_EMAILS": "",
                                   "PRODUCTOS_DATA_DIR": tempfile.mkdtemp(prefix="e2e-productos-"),
                                   "ODOO_URL": "https://ejemplo.odoo.com", "ODOO_DB": "test",
                                   "ODOO_USER": "test", "ODOO_PASSWORD": "test"})]
for p in parches: p.start()

servidor = uvicorn.Server(uvicorn.Config(main.app, host="127.0.0.1", port=PORT, log_level="error"))
threading.Thread(target=servidor.run, daemon=True).start()
for _ in range(100):
    if servidor.started: break
    time.sleep(.1)

TSV = "\t".join(["Referencia interna (SKU)", "Nombre", "Categoría", "Unidad",
                 "Unidades por caja", "Precio de venta", "Costo"]) + "\n" + "\n".join([
    "\t".join(["400001", "Pan Molde Clásico x500g", "Panificados", "Unidades", "12", "2890", "1740"]),
    "\t".join(["400005", "Pan de Campo x600g", "Panificados", "Unidades", "10", "3480", "1990"]),
    "\t".join(["400099", "Pan Nuevo Integral", "Panificados", "Unidades", "8", "1990", "1200"]),
    "\t".join(["400098", "Pan Roto", "Panificado", "Unidades", "6", "1500", "900"]),
])

SALIDA = tempfile.mkdtemp(prefix="e2e-capturas-")
def volver_a_la_entrada(pg):
    """Tras revertir, la pagina se recarga y ya esta en la vista de entrada;
    tras previsualizar, hay que salir con el boton."""
    if pg.is_visible("#btn-otro-archivo"):
        pg.click("#btn-otro-archivo")
    pg.wait_for_selector("#pegar-productos", state="visible", timeout=10000)
    pg.wait_for_timeout(200)


fallas = []
def check(ok, msg):
    print(("  OK   " if ok else "  FALLA ") + msg)
    if not ok: fallas.append(msg)

with sync_playwright() as pw:
    b = pw.chromium.launch(**({"executable_path": os.environ["CHROMIUM"]} if os.getenv("CHROMIUM") else {}))
    ctx = b.new_context(viewport={"width": 1500, "height": 980})
    ctx.add_cookies([{"name": "CF_Authorization", "value": "token-de-prueba",
                      "domain": "127.0.0.1", "path": "/"}])
    pg = ctx.new_page()
    errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(f"http://127.0.0.1:{PORT}/productos"); pg.wait_for_timeout(900)

    print("\n1) La pantalla se sirve y habla con Odoo")
    check("Conectado a Odoo" in pg.text_content("#conexion"), "el status dice que hay conexión")
    cats = pg.evaluate("ESTADO.catalogo.categorias.map(c=>c.name)")
    check("Panificados" in cats and "Galletitas" in cats, f"el catálogo llegó del backend: {cats}")

    print("\n2) Pegar desde Excel y previsualizar")
    pg.fill("#pegar-productos", TSV)
    pg.click("#btn-leer-productos")
    pg.wait_for_selector("#tabla-productos tbody tr", timeout=15000); pg.wait_for_timeout(500)
    estados = pg.evaluate("ESTADO.filas.productos.map(f=>[f.sku,f.estado])")
    esperado = [["400001","igual"],["400005","modificacion"],["400099","alta"],["400098","error"]]
    check(estados == esperado, f"estados: {estados}")
    check(pg.eval_on_selector_all(".viejo","e=>e.length") > 0, "se ve el valor viejo tachado")
    err = pg.text_content("#tabla-productos .err")
    check("¿Quisiste decir: Panificados?" in err, f"sugerencia de categoría: {err.strip()[:70]}")
    check(pg.eval_on_selector("#btn-aplicar","e=>e.disabled"), "aplicar bloqueado por la fila con error")
    check(o.escrituras == [], "NADA se escribió en Odoo durante la previsualización")
    pg.screenshot(path=os.path.join(SALIDA, "e2e-1-preview.png"))

    print("\n3) Corregir el error en la tabla y volver a previsualizar")
    sel = pg.query_selector_all('#tabla-productos tbody tr')[3].query_selector('[data-col="Categoría"] .celda-edit')
    sel.select_option("Panificados"); pg.wait_for_timeout(900)
    estados = pg.evaluate("ESTADO.filas.productos.map(f=>[f.sku,f.estado])")
    check([e[1] for e in estados] == ["igual","modificacion","alta","alta"], f"estados: {estados}")
    check(not pg.eval_on_selector("#btn-aplicar","e=>e.disabled"), "aplicar habilitado")
    check(o.escrituras == [], "sigue sin escribirse nada")

    print("\n4) Aplicar de verdad")
    antes = len(o.datos["product.template"])
    pg.click("#btn-aplicar")
    pg.wait_for_selector("#btn-revertir", timeout=20000); pg.wait_for_timeout(400)
    creados = [r for r in o.datos["product.template"].values() if r.get("default_code") in ("400099","400098")]
    check(len(creados) == 2, f"se crearon 2 productos (había {antes}, hay {len(o.datos['product.template'])})")
    check(o.datos["product.template"][d["p2"]]["list_price"] == 3480.0, "400005 quedó en 3480")
    check(o.datos["product.template"][d["p1"]]["list_price"] == 2890.0, "400001 no se tocó")
    tocados = [a[2][0] for a in o.escrituras if a[1] == "write"]
    check([d["p1"]] not in tocados, "no hubo ningún write sobre la fila sin cambios")
    check(all(r["type"] == "consu" and r["is_storable"] for r in creados), "los fijos de Kairon se aplicaron")
    pg.screenshot(path=os.path.join(SALIDA, "e2e-2-aplicado.png"))

    print("\n5) Revertir la corrida")
    pg.on("dialog", lambda dlg: dlg.accept())
    pg.click("#btn-revertir"); pg.wait_for_timeout(2500)
    check(o.datos["product.template"][d["p2"]]["list_price"] == 3180.0, "400005 volvió a 3180")
    check(all(not r["active"] for r in creados), "los productos creados quedaron archivados, no borrados")
    check(all(r.get("default_code") in ("400099","400098") for r in creados), "siguen existiendo en la base")

    print("\n6) Descargar la plantilla")
    import urllib.request
    pedido = urllib.request.Request(f"http://127.0.0.1:{PORT}/api/productos/plantilla.xlsx")
    pedido.add_header("Cookie", "CF_Authorization=token-de-prueba")
    xlsx = urllib.request.urlopen(pedido).read()
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(xlsx))
    check(set(wb.sheetnames) == {"Productos", "Proveedores", "Precios", "Valores válidos"},
          f"la plantilla trae las cuatro hojas: {wb.sheetnames}")
    validos = [f[1].value for f in wb["Valores válidos"].iter_rows(min_row=2)]
    check("Panificados" in validos and "Mayorista" in validos,
          "trae los valores válidos que Odoo acepta hoy")

    print("\n7) Subir un .xlsx de verdad")
    from openpyxl import Workbook
    wb2 = Workbook(); wb2.remove(wb2.active)
    ws = wb2.create_sheet("Productos")
    ws.append(["Referencia interna (SKU)", "Nombre", "Precio de venta"])
    ws.append(["400005", "Pan de Campo x600g", 3999])
    ws2 = wb2.create_sheet("Precios")
    ws2.append(["SKU (igual al de Productos)", "Lista de precios", "Precio fijo"])
    ws2.append(["400005", "Mayorista", 3500])
    ruta = os.path.join(SALIDA, "plantilla.xlsx"); wb2.save(ruta)

    volver_a_la_entrada(pg)
    # A esta altura ya hubo escrituras legitimas (aplicar y revertir): lo que
    # se comprueba es que estos pasos no agreguen ninguna.
    escrituras_antes = len(o.escrituras)
    pg.set_input_files("#file-productos", ruta)
    pg.click("#btn-leer-productos")
    pg.wait_for_selector("#tabla-productos tbody tr", timeout=15000); pg.wait_for_timeout(400)
    check(pg.evaluate("ESTADO.filas.productos.map(f=>f.estado)") == ["modificacion"],
          "la hoja Productos del .xlsx se leyó y previsualizó")
    check(pg.evaluate("ESTADO.filas.precios.length") == 1,
          "la hoja Precios del mismo archivo también se leyó")
    check(len(o.escrituras) == escrituras_antes, "leer un archivo no escribe nada")

    print("\n8) Modificación masiva sin archivo")
    volver_a_la_entrada(pg)
    escrituras_antes = len(o.escrituras)
    pg.select_option("#m-categoria", "Panificados")
    pg.select_option("#m-columna", "Precio de venta")
    pg.select_option("#m-operacion", "multiplicar")
    pg.fill("#m-valor", "1.15")
    pg.click("#btn-masiva")
    pg.wait_for_selector("#tabla-productos tbody tr", timeout=15000); pg.wait_for_timeout(500)
    estados = pg.evaluate("ESTADO.filas.productos.map(f=>[f.sku,f.estado])")
    check(sorted(e[0] for e in estados) == ["400001", "400005"],
          f"solo los que cambian: {estados}")
    check(all(e[1] == "modificacion" for e in estados), "todos son modificaciones")
    check(len(o.escrituras) == escrituras_antes,
          "la modificación masiva tampoco escribe al previsualizar")
    pg.screenshot(path=os.path.join(SALIDA, "e2e-3-masiva.png"))

    pg.click("#btn-aplicar")
    pg.wait_for_selector("#btn-revertir", timeout=20000); pg.wait_for_timeout(400)
    check(abs(o.datos["product.template"][d["p1"]]["list_price"] - 2890 * 1.15) < 0.01,
          f"400001 quedó en {o.datos['product.template'][d['p1']]['list_price']} (2890 x 1.15)")
    check(abs(o.datos["product.template"][d["p2"]]["list_price"] - 3180 * 1.15) < 0.01,
          f"400005 quedó en {o.datos['product.template'][d['p2']]['list_price']} (3180 x 1.15)")
    check("image_1920" not in [k for m in o.escrituras if m[1] == "write" for k in m[2][1]],
          "no se tocó ningún campo fuera del pedido")

    print("\n9) Errores de JavaScript")
    reales = [e for e in errs if "ERR_CONNECTION" not in e]
    check(not reales, f"sin errores de JS ({reales})")
    b.close()

servidor.should_exit = True
for p in parches: p.stop()
print("\n" + ("="*58) + f"\n{'TODO OK' if not fallas else str(len(fallas)) + ' FALLA(S): ' + '; '.join(fallas)}\n" + "="*58)
sys.exit(1 if fallas else 0)
